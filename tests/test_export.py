from pubpulse.export import authorship_rows, paper_rows, write_xlsx


def _pubs():
    return {"since_year": 2024, "generated": "2026-10-23T00:00:00+00:00",
            "people": [{"slug": "a", "name": "Ann A", "active": True, "orcid": "0000-0001-0000-0001", "openalex_ids": ["A1"], "match": "orcid", "needs_review": False},
                       {"slug": "b", "name": "Bob B", "active": False, "orcid": None, "openalex_ids": [], "match": "name", "needs_review": True}],
            "works": [{"id": "W1", "doi": "https://doi.org/10.1/x", "title": "T", "journal": "J", "publication_date": "2026-01-02", "year": 2026, "type": "article",
                       "authors": ["Ann A", "X", "Bob B"], "author_count": 3, "author_members": ["a", None, "b"], "author_positions": [1, 2, 3],
                       "group_members": ["a", "b"], "counted_members": ["a"], "affiliation_status": {"a": "verified", "b": "mismatch"},
                       "member_affiliations": {"a": ["Maastro"], "b": ["Elsewhere"]}, "verified": True, "counted": True, "approved_external": False,
                       "is_oa": True, "oa_url": "u", "license": "cc-by", "cited_by_count": 2, "first_seen": "2026-10-21", "openalex_url": "o", "landing_url": None, "abstract": "abs"}]}


def test_rows_and_xlsx(tmp_path):
    p = paper_rows(_pubs())[0]
    assert p["group_author_positions"] == "Ann A 1/3; Bob B 3/3" and p["first_author_group"] and p["last_author_group"] and p["counted_team_window"]
    a = authorship_rows(_pubs())
    assert [r["counted_for_member"] for r in a] == [True, False] and a[0]["fraction"] == 0.3333
    out = write_xlsx(_pubs(), {"group_name": "G", "site_url": "s"}, tmp_path / "x.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["Publications", "Group authorships", "People", "Notes"]
    assert wb["Publications"].max_row == 2 and wb["Group authorships"].max_row == 3 and wb["People"].max_row == 3


def test_csv_companion(tmp_path):
    from pubpulse.export import write_csv
    import csv
    out = write_csv(_pubs(), tmp_path / "p.csv")
    rows = list(csv.reader(open(out, encoding="utf-8")))
    assert rows[0][0] == "OpenAlex work id" and rows[1][0] == "W1" and len(rows) == 2
