"""Excel export of everything needed to reproduce the counting: one row per paper, one row per group authorship,
the people list, and a notes sheet with the rules. Generated at every build into dist/data/."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

PAPER_COLUMNS = [
    ("openalex_id", "OpenAlex work id"), ("doi", "DOI"), ("title", "Title"), ("journal", "Journal / source"),
    ("publication_date", "Publication date"), ("year", "Year"), ("type", "Type (OpenAlex)"), ("is_preprint", "Preprint"),
    ("counted_team_window", "Counted (team-window rule)"), ("verified_affiliation", "Verified group affiliation"),
    ("manual_decision", "Manual decision (overrides.yml)"),
    ("group_author_count", "Number of group authors"), ("group_authors", "Group authors"),
    ("group_author_positions", "Group author positions (pos/total)"), ("first_author_group", "First author is group member"),
    ("last_author_group", "Last author is group member"), ("group_affiliation_status", "Affiliation status per group author"),
    ("author_count", "Total authors"), ("all_authors", "All authors (OpenAlex order)"),
    ("open_access", "Open access"), ("oa_url", "Open-access link"), ("license", "Licence"), ("cited_by_count", "Citations (OpenAlex)"),
    ("journal_impact", "Journal impact (OpenAlex 2-yr mean citedness)"),
    ("first_seen", "First seen by pubpulse"), ("openalex_url", "OpenAlex URL"), ("landing_url", "Landing page"), ("abstract", "Abstract"),
]


def paper_rows(pubs: dict) -> list[dict]:
    names = {p["slug"]: p["name"] for p in pubs.get("people", [])}
    rows = []
    for w in pubs.get("works", []):
        members = w.get("group_members") or []
        am = w.get("author_members") or []
        pos = {am[i]: (w.get("author_positions") or [])[i] if i < len(w.get("author_positions") or []) else i + 1 for i in range(len(am)) if am[i]}
        total = w.get("author_count") or len(w.get("authors") or [])
        st = w.get("affiliation_status") or {}
        rows.append({
            "openalex_id": w["id"], "doi": w.get("doi"), "title": w.get("title"), "journal": w.get("journal"),
            "publication_date": w.get("publication_date"), "year": w.get("year"), "type": w.get("type"),
            "is_preprint": w.get("type") == "preprint",
            "counted_team_window": bool(w.get("counted", w.get("verified"))), "verified_affiliation": bool(w.get("verified")),
            "manual_decision": w.get("manual") or "",
            "group_author_count": len(members), "group_authors": "; ".join(names.get(s, s) for s in members),
            "group_author_positions": "; ".join(f"{names.get(s, s)} {pos.get(s, '?')}/{total}" for s in members),
            "first_author_group": bool(am) and am[0] is not None,
            "last_author_group": bool(am) and len(am) == total and am[-1] is not None,
            "group_affiliation_status": "; ".join(f"{names.get(s, s)}: {st.get(s, 'unknown')}" for s in members),
            "author_count": total, "all_authors": "; ".join(w.get("authors") or []),
            "open_access": bool(w.get("is_oa")), "oa_url": w.get("oa_url"), "license": w.get("license"),
            "cited_by_count": w.get("cited_by_count"), "journal_impact": w.get("journal_impact"), "first_seen": w.get("first_seen"),
            "openalex_url": w.get("openalex_url"), "landing_url": w.get("landing_url"), "abstract": w.get("abstract"),
        })
    return rows


AUTHORSHIP_COLUMNS = [
    ("openalex_id", "OpenAlex work id"), ("doi", "DOI"), ("title", "Title"), ("year", "Year"), ("type", "Type"),
    ("member", "Group author"), ("slug", "Member id"), ("position", "Author position"), ("author_count", "Total authors"),
    ("fraction", "Fractional credit (1/total)"), ("affiliation_status", "Affiliation status"),
    ("affiliations", "Affiliations recorded by OpenAlex"), ("counted_for_member", "Counted for this member"),
    ("paper_counted", "Paper counted (team-window)"), ("member_active", "Member currently in team"),
]


def authorship_rows(pubs: dict) -> list[dict]:
    people = {p["slug"]: p for p in pubs.get("people", [])}
    rows = []
    for w in pubs.get("works", []):
        am = w.get("author_members") or []
        positions = w.get("author_positions") or []
        total = w.get("author_count") or len(w.get("authors") or [])
        pos = {am[i]: positions[i] if i < len(positions) else i + 1 for i in range(len(am)) if am[i]}
        counted = set(w.get("counted_members") or ([s for s in w.get("group_members", [])] if w.get("counted") else []))
        for s in w.get("group_members") or []:
            p = people.get(s, {})
            rows.append({
                "openalex_id": w["id"], "doi": w.get("doi"), "title": w.get("title"), "year": w.get("year"), "type": w.get("type"),
                "member": p.get("name", s), "slug": s, "position": pos.get(s), "author_count": total,
                "fraction": round(1 / total, 4) if total else None,
                "affiliation_status": (w.get("affiliation_status") or {}).get(s, "unknown"),
                "affiliations": "; ".join((w.get("member_affiliations") or {}).get(s, [])),
                "counted_for_member": s in counted, "paper_counted": bool(w.get("counted", w.get("verified"))),
                "member_active": bool(p.get("active")),
            })
    return rows


PEOPLE_COLUMNS = [
    ("slug", "Member id"), ("name", "Name"), ("role", "Role (staff page)"), ("employer", "Employer (staff page)"),
    ("active", "Currently in team"), ("orcid", "ORCID"), ("openalex_ids", "OpenAlex author ids"), ("match", "Identity basis"),
    ("needs_review", "Needs review (no identifier)"), ("profile_url", "Staff profile"),
]


def people_rows(pubs: dict) -> list[dict]:
    return [{
        "slug": p["slug"], "name": p["name"], "role": p.get("role"), "employer": p.get("employer"), "active": bool(p.get("active")),
        "orcid": p.get("orcid"), "openalex_ids": "; ".join(p.get("openalex_ids") or []), "match": p.get("match"),
        "needs_review": bool(p.get("needs_review")), "profile_url": p.get("profile_url"),
    } for p in pubs.get("people", [])]


def write_xlsx(pubs: dict, cfg: dict, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def sheet(ws, columns, rows, widths=None):
        ws.append([label for _, label in columns])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append([r.get(k) for k, _ in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, (k, _) in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(k, 16)

    ws = wb.active
    ws.title = "Publications"
    sheet(ws, PAPER_COLUMNS, paper_rows(pubs), {"title": 60, "journal": 30, "doi": 34, "group_authors": 36, "group_author_positions": 40,
                                                  "group_affiliation_status": 40, "all_authors": 60, "abstract": 80, "oa_url": 40, "openalex_url": 34})
    sheet(wb.create_sheet("Group authorships"), AUTHORSHIP_COLUMNS, authorship_rows(pubs), {"title": 60, "doi": 34, "member": 24, "affiliations": 60})
    sheet(wb.create_sheet("People"), PEOPLE_COLUMNS, people_rows(pubs), {"name": 26, "role": 28, "employer": 36, "openalex_ids": 28, "profile_url": 44})

    n = wb.create_sheet("Notes")
    works = pubs.get("works", [])
    counted = [w for w in works if w.get("counted", w.get("verified"))]
    for line in [
        f"pubpulse export · generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} · {cfg.get('group_name', '')}",
        f"Source: OpenAlex (CC0). Coverage: publications from {pubs.get('since_year')} onwards. Data as of {pubs.get('generated', '')}.",
        "",
        "Team-window rule (column 'Counted (team-window rule)'): a paper counts when a group author has a verified group affiliation, or when a "
        "group member was in the team at publication time and OpenAlex lists no other affiliation for them (none, or only a bare Maastricht University / MUMC+ / Radiation Oncology). "
        "Listed with another institution only: "
        "not counted. Outside the team window, and for people matched by name only, a verified group affiliation is required.",
        "Verified group affiliation: the authorship carries Maastro, Medical Data Works, Clinical Data Science, or Radiation Oncology/GROW together with Maastricht/MUMC.",
        "Affiliation status per author: verified / unspecific (only a bare Maastricht University, MUMC+ or Radiation Oncology) / mismatch (another department or institution recorded) / unknown (OpenAlex has no affiliation for this authorship).",
        "Manual decisions (overrides.yml): confirmed = counted by hand; rejected papers are not in this file.",
        "Preprints are flagged in 'Preprint' and hidden on the website by default; they are included here.",
        "Sheet 'Group authorships' has one row per (paper, group author) for fractional or per-person counting; 'Fractional credit' = 1 / total authors.",
        "",
        f"Totals: {len(works)} papers in the archive; {len(counted)} counted; {sum(1 for w in counted if w.get('type') != 'preprint')} counted excluding preprints; "
        f"{sum(1 for w in works if w.get('verified'))} with a verified affiliation.",
        f"Website: {cfg.get('site_url', '')}",
    ]:
        n.append([line])
    n.column_dimensions["A"].width = 140
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_csv(pubs: dict, path: Path) -> Path:
    """The Publications sheet as UTF-8 CSV (companion of the xlsx; read by the review sheet script)."""
    import csv

    rows = paper_rows(pubs)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([label for _, label in PAPER_COLUMNS])
        for r in rows:
            w.writerow([("" if r.get(k) is None else r.get(k)) for k, _ in PAPER_COLUMNS])
    return path
